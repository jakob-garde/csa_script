"""Structure of the order list, checked against a tiny workbook."""

import unittest
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import Workbook

from csa.html import render_html, render_index
from csa.pages import family_pages, group_pages
from csa.picking import picking_tables
from csa.read import read_week


def _save(workbook):
    directory = TemporaryDirectory()
    path = Path(directory.name) / "week.xlsx"
    workbook.save(path)
    return directory, path


def _group(workbook, title, families, rows):
    if workbook.sheetnames == ["Sheet"] and workbook["Sheet"]["A2"].value is None:
        sheet = workbook["Sheet"]
        sheet.title = title
    else:
        sheet = workbook.create_sheet(title)
    sheet["A2"] = "öppen"
    sheet["B2"] = "Total"
    for index, name in enumerate(families):
        sheet.cell(2, 3 + index * 2, name)
        sheet.cell(2, 4 + index * 2, "text")
    for row_index, row in enumerate(rows, start=3):
        for column, value in enumerate(row, start=1):
            if value is not None:
                sheet.cell(row_index, column, value)
    return sheet


class ReadStructureTest(unittest.TestCase):
    def test_families_are_amount_and_note_pairs(self):
        workbook = Workbook()
        _group(
            workbook,
            'grupp "Annica"',
            ["Andreas", "Annica"],
            [
                ["Morot", 3, 2, "tvätta", 1, None],
                ["Mejeri", None],
                ["Köttfärs", 2, 2, None, None, None],
                ["Små Plantor", None],
            ],
        )
        directory, path = _save(workbook)
        try:
            week = read_week(path)
        finally:
            directory.cleanup()

        group = week.groups[0]
        self.assertEqual(group.clean_name, "Annica")
        self.assertEqual(group.families, ("Andreas", "Annica"))
        morot = group.items[0]
        self.assertEqual(morot.name, "Morot")
        self.assertEqual(morot.amounts, (2, 1))
        self.assertEqual(morot.notes, ("tvätta", None))

    def test_picking_list_lines_up_items_by_name(self):
        workbook = Workbook()
        _group(
            workbook,
            "Grupp Maria",
            ["A"],
            [
                ["Morot", 2, 2],
                ["Mejeri", None],
                ["Köttfärs", 1, 1],
                ["Små Plantor", None],
            ],
        )
        _group(
            workbook,
            "grupp Mattias",
            ["B"],
            [
                ["Morot", 3, 3],
                ["Mejeri", None],
                ["Köttfärs", 4, 4],
                ["Små Plantor", None],
            ],
        )
        directory, path = _save(workbook)
        try:
            tables = picking_tables(read_week(path))
        finally:
            directory.cleanup()

        by_label = dict(tables)
        self.assertEqual(by_label["Veggies"][0], ["", "Total", "Group Maria", "Group Mattias"])
        self.assertEqual(by_label["Veggies"][1], ["Morot", "5", "2", "3"])
        self.assertEqual(by_label["Dairy & Meat"][1][0], "Mejeri")
        meat = {row[0]: row for row in by_label["Dairy & Meat"][1:]}
        self.assertEqual(meat["Köttfärs"], ["Köttfärs", "5", "1", "4"])

    def test_long_notice_without_a_number_is_left_out(self):
        notice = "Detta är en lång notis till gruppen " + ("x" * 60)
        self.assertGreater(len(notice), 80)
        workbook = Workbook()
        _group(
            workbook,
            "Grupp Maria",
            ["A"],
            [
                ["Morot", 1, 1],
                [notice, None],
                ["Mejeri", None],
                ["Små Plantor", None],
            ],
        )
        directory, path = _save(workbook)
        try:
            veggies = dict(picking_tables(read_week(path)))["Veggies"]
        finally:
            directory.cleanup()

        names = [row[0] for row in veggies]
        self.assertIn("Morot", names)
        self.assertNotIn(notice, names)

    def test_a_formula_keeps_a_row_even_when_the_total_is_blank(self):
        workbook = Workbook()
        _group(
            workbook,
            "Grupp Maria",
            ["A"],
            [
                ["Pumpa", "=SUM(C3:C3)", None],
                ["Mejeri", None],
                ["Små Plantor", None],
            ],
        )
        directory, path = _save(workbook)
        try:
            veggies = dict(picking_tables(read_week(path)))["Veggies"]
        finally:
            directory.cleanup()

        self.assertEqual(veggies[1][0], "Pumpa")
        self.assertEqual(veggies[1][1], "")

    def test_season_ended_band_is_left_out(self):
        workbook = Workbook()
        _group(
            workbook,
            "Grupp Maria",
            ["A"],
            [
                ["Morot", 1, 1],
                ["Mejeri", None],
                ["Köttfärs", 2, 2],
                ["Små Plantor", None],
                ["Slut för säsongen", None],
            ],
        )
        directory, path = _save(workbook)
        try:
            labels = [label for label, _rows in picking_tables(read_week(path))]
        finally:
            directory.cleanup()

        self.assertEqual(labels, ["Veggies", "Dairy & Meat"])

    def test_words_inside_a_product_name_do_not_end_the_band(self):
        workbook = Workbook()
        _group(
            workbook,
            "Grupp Maria",
            ["A"],
            [
                ["Sallad Isberg. SLUT för säsongen.", 2, 2],
                ["Mejeri", None],
                ["Små Plantor", None],
            ],
        )
        directory, path = _save(workbook)
        try:
            veggies = dict(picking_tables(read_week(path)))["Veggies"]
        finally:
            directory.cleanup()

        self.assertEqual(veggies[1][0], "Sallad Isberg. SLUT för säsongen.")


class FamilyPageTest(unittest.TestCase):
    def _week(self):
        workbook = Workbook()
        sheet = _group(
            workbook,
            'grupp "Annica"',
            ["Andreas", "Tyst"],
            [
                ["Morot", 2, 2, "tvätta", None, None],
                ["Mejeri", None],
                ["Köttfärs", 1, 1, None, None, None],
                ["Små Plantor", None],
                ["Slut för säsongen", None],
            ],
        )
        sheet["A1"] = "Hos Annica, torsdag."
        return _save(workbook)

    def test_one_page_per_family_with_the_bands_together(self):
        directory, path = self._week()
        try:
            pages = list(family_pages(read_week(path)))
        finally:
            directory.cleanup()

        self.assertEqual(len(pages), 2)
        rel, title, intro, sections = pages[0]
        self.assertEqual(rel, "families/group_annica/andreas.html")
        self.assertEqual(title, "Group Annica — Andreas")
        self.assertEqual(pages[1][0], "families/group_annica/tyst.html")
        self.assertEqual(
            pages[1][3],
            [
                ("Veggies", [["", "Antal", "text"], ["Morot", "", ""]]),
                ("Dairy & Meat", [["", "Antal", "text"], ["Köttfärs", "", ""]]),
            ],
        )
        self.assertEqual(intro, "Hos Annica, torsdag.")
        labels = [label for label, _rows in sections]
        self.assertEqual(labels, ["Veggies", "Dairy & Meat"])
        veggies = dict(sections)["Veggies"]
        self.assertEqual(veggies[1], ["Morot", "2", "tvätta"])

    def test_group_totals_are_that_groups_column(self):
        workbook = Workbook()
        _group(workbook, "Grupp Maria", ["A"], [
            ["Morot", 2, 2],
            ["Mejeri", None],
            ["Små Plantor", None],
        ])
        _group(workbook, "grupp Mattias", ["B"], [
            ["Morot", 3, 3],
            ["Mejeri", None],
            ["Små Plantor", None],
        ])
        directory, path = _save(workbook)
        try:
            pages = list(group_pages(read_week(path)))
        finally:
            directory.cleanup()

        self.assertEqual(pages[0][0], "groups/group_maria.html")
        self.assertEqual(pages[0][1], "Group Maria")
        self.assertEqual(pages[1][0], "groups/group_mattias.html")
        self.assertEqual(pages[1][1], "Group Mattias")
        maria = dict(pages[0][3])["Veggies"]
        self.assertEqual(maria[1], ["Morot", "2"])


class PagePasswordTest(unittest.TestCase):
    def test_every_page_asks_before_showing_the_lists(self):
        page = render_html("Picking list", [("Veggies", [["", "Total"], ["Morot", "2"]])])
        index = render_index(
            "This week",
            [("Group Annica", "groups/group_annica.html", [])],
            "2026-10-25 11:51",
        )
        for text in (page, index):
            self.assertLess(text.index('id="gate"'), text.index('id="page"'))
            self.assertLess(text.index('id="page"'), text.index("<h1>"))
            self.assertIn("=== 'tomat'", text)
            self.assertIn("localStorage", text)
            self.assertIn("#page { display: none; }", text)
        self.assertLess(index.index("<h1>This week</h1>"), index.index("2026-10-25 11:51"))
        self.assertLess(index.index("2026-10-25 11:51"), index.index("Picking list"))


if __name__ == "__main__":
    unittest.main()
