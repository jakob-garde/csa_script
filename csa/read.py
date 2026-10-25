"""Read an order-list export into a Week."""

from openpyxl import load_workbook

from csa.model import Group, Item, Week

MEJERI = "Mejeri"
PLANTOR = "Små Plantor"


def read_week(path):
    values = load_workbook(path, data_only=True)
    formulas = load_workbook(path, data_only=False)
    groups = tuple(
        _read_group(values[name], formulas[name])
        for name in values.sheetnames
    )
    values.close()
    formulas.close()
    return Week(groups)


def _read_group(values, formulas):
    families = _families(values)
    last_row = max(values.max_row or 1, formulas.max_row or 1)
    items = []
    for row in range(3, last_row + 1):
        name = _text(formulas.cell(row, 1).value)
        if name == "":
            name = _text(values.cell(row, 1).value)
        formula = formulas.cell(row, 2).value
        has_formula = isinstance(formula, str) and formula.startswith("=")
        total = values.cell(row, 2).value
        if total is None and not has_formula:
            total = formula
        amounts = []
        notes = []
        for index in range(len(families)):
            column = 3 + index * 2
            amounts.append(_cell(values, formulas, row, column))
            notes.append(_cell(values, formulas, row, column + 1))
        if name == "" and total is None and not any(amounts) and not any(notes):
            continue
        items.append(Item(name, total, has_formula, tuple(amounts), tuple(notes)))
    note = _text(values.cell(1, 1).value)
    if note == "":
        note = _text(formulas.cell(1, 1).value)
    return Group(values.title, families, tuple(items), note)


def _families(sheet):
    families = []
    column = 3
    while column <= (sheet.max_column or 0):
        name = sheet.cell(2, column).value
        note_header = sheet.cell(2, column + 1).value
        if name is None and note_header is None:
            break
        families.append(_text(name))
        column += 2
    return tuple(families)


def _cell(values, formulas, row, column):
    raw = values.cell(row, column).value
    if raw is None:
        other = formulas.cell(row, column).value
        if not (isinstance(other, str) and other.startswith("=")):
            raw = other
    return raw


def _text(value):
    if value is None:
        return ""
    return str(value).strip()
